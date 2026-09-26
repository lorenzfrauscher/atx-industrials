#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
auswertung_verifikation.py -- wertet die Handpruefung aus Schritt 3 aus

Liest pruefliste.xlsx (die Arbeitsdatei) und erzeugt die Ergebnistabellen fuer
verification.md. Schreibt ausserdem pruefliste.csv fuer das Repository, damit die
Rohurteile ohne Excel lesbar sind.

Aufruf:  python3 auswertung_verifikation.py
"""

import collections
import csv
import io
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
FELDER = ["strategische_prioritaeten", "wachstumsmaerkte",
          "hauptrisiken", "capex_ma_signale"]
KLASSEN = ["korrekt", "beleg_stuetzt_nicht", "inhalt_verzerrt",
           "kategorie_falsch", "seitenangabe_falsch", "zitat_erfunden",
           "kontext_entrissen", "unklar"]


def laden(pfad):
    import openpyxl
    ws = openpyxl.load_workbook(pfad, data_only=True).active
    kopf = [str(c.value or "").strip().lower() for c in ws[1]]
    zeilen = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        zeilen.append(dict(zip(kopf, [("" if v is None else v) for v in r])))
    return zeilen


def urteile(z):
    """Trennt an Komma UND Semikolon. Beim Eintragen von Hand entstehen beide
    Varianten, und ein nicht erkannter Trenner erzeugt eine Phantomklasse."""
    roh = str(z.get("urteil", "")).replace(";", ",")
    gefunden = [x.strip() for x in roh.split(",") if x.strip()]
    unbekannt = [x for x in gefunden if x not in KLASSEN]
    if unbekannt:
        print("   Hinweis: unbekanntes Urteil %s bei %s" % (unbekannt, z.get("id")))
    return gefunden


def tab(kopf, zeilen):
    aus = ["| " + " | ".join(kopf) + " |",
           "|" + "|".join(["---"] * len(kopf)) + "|"]
    for z in zeilen:
        aus.append("| " + " | ".join(str(x) for x in z) + " |")
    return "\n".join(aus)


def main():
    pfad = os.path.join(HERE, "pruefliste.xlsx")
    if not os.path.exists(pfad):
        sys.exit("pruefliste.xlsx nicht gefunden.")
    zeilen = laden(pfad)
    offen = [z["id"] for z in zeilen if not urteile(z)]
    if offen:
        print("WARNUNG: %d Eintraege ohne Urteil: %s" % (len(offen), offen[:10]))

    geprueft = [z for z in zeilen if urteile(z)]
    n = len(geprueft)
    ok = sum(1 for z in geprueft if urteile(z) == ["korrekt"])

    t = ["\n## 5. Ergebnisse der Handpruefung\n",
         "Gepruefte Elemente: **%d** aus drei zufaellig gezogenen Unternehmen. "
         "Ein Element kann mehrere Fehlerklassen tragen, deshalb uebersteigt die "
         "Summe der Klassen die Zahl der fehlerhaften Elemente.\n" % n,
         "**Vollstaendig korrekt: %d von %d, also %.0f Prozent.**\n" % (ok, n, 100.0 * ok / n)]

    # Je Unternehmen
    t.append("\n### Nach Unternehmen\n")
    je = collections.defaultdict(lambda: collections.Counter())
    for z in geprueft:
        f = str(z["unternehmen"])
        je[f]["n"] += 1
        if urteile(z) == ["korrekt"]:
            je[f]["ok"] += 1
        for u in urteile(z):
            if u != "korrekt":
                je[f][u] += 1
    rows = []
    for f, v in sorted(je.items()):
        rest = ", ".join("%s %d" % (k, x) for k, x in v.items() if k not in ("n", "ok")) or "-"
        rows.append([f, v["n"], "%d (%.0f%%)" % (v["ok"], 100.0 * v["ok"] / v["n"]), rest])
    t.append(tab(["Unternehmen", "geprueft", "korrekt", "Fehlerklassen"], rows))

    # Je Feldtyp
    t.append("\n\n### Nach Feldtyp\n")
    jf = collections.defaultdict(lambda: collections.Counter())
    for z in geprueft:
        f = str(z["feld"])
        jf[f]["n"] += 1
        if urteile(z) == ["korrekt"]:
            jf[f]["ok"] += 1
        for u in urteile(z):
            if u != "korrekt":
                jf[f][u] += 1
    rows = []
    for f in FELDER:
        v = jf.get(f)
        if not v:
            continue
        rest = ", ".join("%s %d" % (k, x) for k, x in v.items() if k not in ("n", "ok")) or "-"
        rows.append([f, v["n"], "%d (%.0f%%)" % (v["ok"], 100.0 * v["ok"] / v["n"]), rest])
    t.append(tab(["Feldtyp", "geprueft", "korrekt", "Fehlerklassen"], rows))

    # Je Herkunftsabschnitt
    t.append("\n\n### Nach Herkunftsabschnitt\n")
    ja = collections.defaultdict(lambda: collections.Counter())
    for z in geprueft:
        a = str(z.get("quelle_abschnitt", "unbekannt"))
        ja[a]["n"] += 1
        if urteile(z) == ["korrekt"]:
            ja[a]["ok"] += 1
    rows = [[a, v["n"], "%d (%.0f%%)" % (v["ok"], 100.0 * v["ok"] / v["n"])]
            for a, v in sorted(ja.items())]
    t.append(tab(["Abschnitt", "geprueft", "korrekt"], rows))

    # Fehlerklassen gesamt
    t.append("\n\n### Fehlerklassen ueber alle geprueften Elemente\n")
    c = collections.Counter()
    for z in geprueft:
        for u in urteile(z):
            if u != "korrekt":
                c[u] += 1
    rows = [[k, v, "%.1f%%" % (100.0 * v / n)] for k, v in c.most_common()]
    t.append(tab(["Fehlerklasse", "Faelle", "Anteil an allen geprueften"], rows))

    # Maschinell vs. von Hand
    t.append("\n\n### Was die Maschine gesehen hat und was nicht\n")
    maschinell = sum(1 for z in geprueft if z.get("treffertyp") == "nicht_gefunden")
    nur_hand = sum(1 for z in geprueft
                   if any(u in ("beleg_stuetzt_nicht", "inhalt_verzerrt",
                                "kategorie_falsch", "kontext_entrissen") for u in urteile(z))
                   and z.get("treffertyp") != "nicht_gefunden")
    t.append(tab(["Befund", "Anzahl"], [
        ["von der Zitatsuche als erfunden gemeldet", maschinell],
        ["nur durch Handpruefung gefunden", nur_hand]]))
    t.append("\nDie automatische Zitatpruefung beweist, dass ein Satz im Dokument steht. "
             "Sie beweist nicht, dass er die daraus abgeleitete Aussage traegt. Diese "
             "Luecke schliesst nur die Handpruefung.\n")

    ziel = os.path.join(HERE, "verifikation_ergebnisse.md")
    io.open(ziel, "w", encoding="utf-8").write("\n".join(t) + "\n")
    print("geschrieben: verifikation_ergebnisse.md")

    # Rohurteile als CSV fuers Repository
    if zeilen:
        with io.open(os.path.join(HERE, "pruefliste.csv"), "w",
                     encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(zeilen[0].keys()))
            w.writeheader()
            w.writerows(zeilen)
        print("geschrieben: pruefliste.csv (%d Zeilen)" % len(zeilen))

    print("\n%d geprueft, %d vollstaendig korrekt (%.0f%%)" % (n, ok, 100.0 * ok / n))
    for k, v in c.most_common():
        print("   %-24s %3d  (%.1f%%)" % (k, v, 100.0 * v / n))


if __name__ == "__main__":
    main()
